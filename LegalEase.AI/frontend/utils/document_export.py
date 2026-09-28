from io import BytesIO

from docx import Document
from fpdf import FPDF


def format_txt(text: str) -> str:
    return text


def format_docx(text: str, document_type: str = "") -> bytes:
    document = Document()
    if document_type:
        document.add_heading(document_type, level=1)

    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def format_pdf(text: str, document_type: str = "") -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=11)

    if document_type:
        safe_title = document_type.encode("latin-1", "replace").decode("latin-1")
        pdf.set_font("Helvetica", style="B", size=16)
        pdf.multi_cell(0, 10, safe_title)
        pdf.ln(4)
        pdf.set_font("Helvetica", size=11)

    for paragraph in text.splitlines():
        safe_paragraph = paragraph.encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 7, safe_paragraph)

    return bytes(pdf.output())